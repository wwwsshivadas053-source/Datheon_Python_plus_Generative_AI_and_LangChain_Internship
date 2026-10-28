import os
from pathlib import Path
from typing import List, Dict, Any
from langchain_core.documents import Document
import pymupdf as fitz  # Modern PyMuPDF API


class DocumentLoaderService:
    """
    Production-grade document loader supporting PDF, TXT, DOCX, and Markdown files.
    Preserves document structure and 1-indexed page metadata for precise citations.
    """

    @staticmethod
    def load_pdf(file_path: str) -> List[Document]:
        """
        Extracts text from PDF page by page using PyMuPDF (fitz) for speed and fidelity.
        Returns a list of LangChain Document objects with 1-indexed page numbers.
        """
        docs: List[Document] = []
        path_obj = Path(file_path)
        file_name = path_obj.name

        try:
            doc = fitz.open(file_path)
            for page_idx in range(len(doc)):
                page = doc[page_idx]
                text = page.get_text("text").strip()
                if text:  # Ignore completely empty pages
                    metadata = {
                        "source": str(file_path),
                        "file_name": file_name,
                        "page": page_idx + 1,  # 1-indexed for human citations
                        "total_pages": len(doc),
                        "char_count": len(text)
                    }
                    docs.append(Document(page_content=text, metadata=metadata))
            doc.close()
        except Exception as e:
            raise RuntimeError(f"Failed to parse PDF document '{file_name}': {str(e)}")

        return docs

    @staticmethod
    def load_txt_or_md(file_path: str) -> List[Document]:
        """Loads plain text or markdown files."""
        path_obj = Path(file_path)
        file_name = path_obj.name

        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read().strip()

            if not content:
                return []

            metadata = {
                "source": str(file_path),
                "file_name": file_name,
                "page": 1,
                "total_pages": 1,
                "char_count": len(content)
            }
            return [Document(page_content=content, metadata=metadata)]
        except Exception as e:
            raise RuntimeError(f"Failed to read text file '{file_name}': {str(e)}")

    @staticmethod
    def load_docx(file_path: str) -> List[Document]:
        """Loads Microsoft Word .docx documents."""
        path_obj = Path(file_path)
        file_name = path_obj.name

        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n\n".join(paragraphs).strip()

            if not content:
                return []

            metadata = {
                "source": str(file_path),
                "file_name": file_name,
                "page": 1,
                "total_pages": 1,
                "char_count": len(content)
            }
            return [Document(page_content=content, metadata=metadata)]
        except Exception as e:
            raise RuntimeError(f"Failed to read DOCX file '{file_name}': {str(e)}")

    @classmethod
    def load_document(cls, file_path: str) -> List[Document]:
        """Dispatches loading according to file extension."""
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            return cls.load_pdf(file_path)
        elif ext in [".txt", ".md"]:
            return cls.load_txt_or_md(file_path)
        elif ext == ".docx":
            return cls.load_docx(file_path)
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Supported: PDF, TXT, MD, DOCX.")

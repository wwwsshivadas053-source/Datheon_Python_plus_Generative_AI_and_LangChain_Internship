"""Document Loader: Multi-format unstructured text ingestion engine.

Supports ingestion from plain text (.txt), Markdown (.md), PDF documents (.pdf),
Word documents (.docx), and direct raw string buffers.
"""

import os
from pathlib import Path
from typing import Union


class DocumentLoader:
    """Unified loader for complex unstructured documents."""

    @classmethod
    def load_from_file(cls, file_path: Union[str, Path]) -> str:
        """Load and extract text from a file path based on its extension."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Document not found at path: {file_path}")

        ext = path.suffix.lower()

        if ext in (".txt", ".md", ".log", ".csv", ".json"):
            return cls._load_plain_text(path)
        elif ext == ".pdf":
            return cls._load_pdf(path)
        elif ext in (".docx", ".doc"):
            return cls._load_docx(path)
        else:
            # Fallback to plain text attempt
            return cls._load_plain_text(path)

    @staticmethod
    def _load_plain_text(path: Path) -> str:
        """Read standard UTF-8 text file with fallback encodings."""
        try:
            return path.read_text(encoding="utf-8").strip()
        except UnicodeDecodeError:
            return path.read_text(encoding="latin-1").strip()

    @staticmethod
    def _load_pdf(path: Path) -> str:
        """Extract text from PDF pages using PyMuPDF (fitz)."""
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(path))
            pages_text = []
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text("text")
                if text.strip():
                    pages_text.append(f"--- PAGE {page_num + 1} ---\n{text.strip()}")
            doc.close()
            return "\n\n".join(pages_text).strip()
        except ImportError:
            raise ImportError(
                "PyMuPDF ('pymupdf') is required to parse PDF files. Install with 'pip install pymupdf'."
            )
        except Exception as e:
            raise RuntimeError(f"Error parsing PDF document {path.name}: {str(e)}")

    @staticmethod
    def _load_docx(path: Path) -> str:
        """Extract text and tables from DOCX files using python-docx."""
        try:
            import docx
            doc = docx.Document(str(path))
            full_text = []

            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())

            # Extract tables as formatted text
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_rows.append(" | ".join(row_data))
                if table_rows:
                    full_text.append("\n[TABLE DATA]\n" + "\n".join(table_rows))

            return "\n\n".join(full_text).strip()
        except ImportError:
            raise ImportError(
                "python-docx ('python-docx') is required to parse DOCX files. Install with 'pip install python-docx'."
            )
        except Exception as e:
            raise RuntimeError(f"Error parsing Word document {path.name}: {str(e)}")

    @classmethod
    def load(cls, source: Union[str, Path]) -> str:
        """Auto-detect whether input is a file path or direct string content."""
        if isinstance(source, Path) or (isinstance(source, str) and (os.path.exists(source) and ("\n" not in source))):
            return cls.load_from_file(source)
        return str(source).strip()
